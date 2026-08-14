"""Safe resume extraction and deterministic first-pass review."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import NamedTuple

from interview_os.core.state import (
    ResumeClaim,
    ResumeFileMetadata,
    ResumeIssueSeverity,
    ResumeReview,
    ResumeValidationIssue,
)

MAX_RESUME_BYTES = 10 * 1024 * 1024
MAX_RESUME_CHARACTERS = 100_000
SUPPORTED_RESUME_TYPES = {".pdf", ".docx"}


class _EmploymentInterval(NamedTuple):
    start_month: int
    end_month: int
    statement: str
    start_is_precise: bool
    end_is_precise: bool


class ResumeProcessingError(ValueError):
    """Raised when an uploaded resume cannot be handled safely."""


class ScannedPDFError(ResumeProcessingError):
    """Signals that a valid PDF needs OCR because it has no text layer."""

    def __init__(self, page_count: int):
        super().__init__("扫描版 PDF 没有文字层")
        self.page_count = page_count


class ResumeProcessor:
    """Extract text, then produce a bounded near-linear deterministic review."""

    def process(self, filename: str, content: bytes) -> tuple[str, ResumeReview]:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_RESUME_TYPES:
            raise ResumeProcessingError("仅支持 PDF 和 Word (.docx) 简历")
        if not content:
            raise ResumeProcessingError("上传的简历为空")
        if len(content) > MAX_RESUME_BYTES:
            raise ResumeProcessingError("简历不能超过 10 MB")

        text, pages = (
            self._extract_pdf(content) if suffix == ".pdf" else self._extract_docx(content)
        )
        had_encoding_artifacts = bool(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufffd]", text))
        text = self._normalize(text)
        if not text:
            if suffix == ".pdf":
                raise ScannedPDFError(pages)
            raise ResumeProcessingError("未能从简历中提取文字")
        if len(text) > MAX_RESUME_CHARACTERS:
            text = text[:MAX_RESUME_CHARACTERS]

        return text, self._build_review(
            filename,
            len(content),
            pages,
            text,
            had_encoding_artifacts=had_encoding_artifacts,
        )

    def process_ocr_text(
        self,
        filename: str,
        content: bytes,
        text: str,
        pages: int,
    ) -> tuple[str, ResumeReview]:
        """Review text produced by local OCR while preserving its uncertainty."""
        normalized = self._normalize(text)
        if len(normalized) < 40:
            raise ResumeProcessingError("自动 OCR 未能提取足够文字，请上传更清晰的扫描件")
        if len(normalized) > MAX_RESUME_CHARACTERS:
            normalized = normalized[:MAX_RESUME_CHARACTERS]
        review = self._build_review(filename, len(content), pages, normalized)
        review.issues.insert(
            0,
            ResumeValidationIssue(
                code="ocr_transcription_unverified",
                severity=ResumeIssueSeverity.WARNING,
                field="document",
                message="扫描件已由本地 OCR 转写，可能存在错字或漏行，请对照原件核对",
            ),
        )
        return normalized, review

    @staticmethod
    def _build_review(
        filename: str,
        size_bytes: int,
        pages: int,
        text: str,
        *,
        had_encoding_artifacts: bool = False,
    ) -> ResumeReview:
        return ResumeReview(
            metadata=ResumeFileMetadata(
                filename=Path(filename).name,
                file_type=Path(filename).suffix.lower().removeprefix("."),
                size_bytes=size_bytes,
                page_count=pages,
                character_count=len(text),
            ),
            issues=ResumeProcessor._find_issues(
                text, had_encoding_artifacts=had_encoding_artifacts
            ),
            claims=ResumeProcessor._find_claims(text),
            reviewed_at=datetime.now(timezone.utc),
        )

    @staticmethod
    def _extract_pdf(content: bytes) -> tuple[str, int]:
        try:
            import pdfplumber

            with pdfplumber.open(BytesIO(content)) as pdf:
                page_texts = [page.extract_text() or "" for page in pdf.pages]
                return "\n\n".join(page_texts), len(pdf.pages)
        except ResumeProcessingError:
            raise
        except Exception as exc:
            raise ResumeProcessingError("PDF 文件损坏或格式不受支持") from exc

    @staticmethod
    def _extract_docx(content: bytes) -> tuple[str, int]:
        try:
            from docx import Document

            document = Document(BytesIO(content))
            blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                row_blocks = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        row_blocks.append(" | ".join(cells))
                if row_blocks:
                    blocks.append("[表格]")
                    blocks.extend(row_blocks)
            return "\n".join(blocks), 1
        except Exception as exc:
            raise ResumeProcessingError("Word 文件损坏或不是有效的 .docx 文件") from exc

    @staticmethod
    def _normalize(text: str) -> str:
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
        return "\n".join(line for line in lines if line).strip()

    @staticmethod
    def _find_issues(
        text: str, *, had_encoding_artifacts: bool = False
    ) -> list[ResumeValidationIssue]:
        issues: list[ResumeValidationIssue] = []
        if had_encoding_artifacts:
            issues.append(
                ResumeValidationIssue(
                    code="encoding_artifacts",
                    severity=ResumeIssueSeverity.WARNING,
                    field="document",
                    message="PDF 包含异常字体编码，个别公司名或符号可能需要对照原文件确认",
                )
            )
        if len(text) < 200:
            issues.append(
                ResumeValidationIssue(
                    code="low_text",
                    severity=ResumeIssueSeverity.WARNING,
                    message="提取到的文字较少，请核对解析结果",
                )
            )
        if not re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text):
            issues.append(
                ResumeValidationIssue(
                    code="missing_email",
                    severity=ResumeIssueSeverity.WARNING,
                    field="contact",
                    message="未识别到邮箱地址",
                )
            )
        if not re.search(r"(?:\+?86[- ]?)?1[3-9]\d{9}|\+?\d[\d ()-]{7,}\d", text):
            issues.append(
                ResumeValidationIssue(
                    code="missing_phone",
                    severity=ResumeIssueSeverity.INFO,
                    field="contact",
                    message="未识别到联系电话",
                )
            )
        section_groups = {
            "experience": ("工作经历", "工作经验", "experience", "employment"),
            "education": ("教育经历", "教育背景", "education", "university"),
            "skills": (
                "技能",
                "专业能力",
                "核心能力",
                "skills",
                "technologies",
                "capabilities",
                "competencies",
                "expertise",
            ),
        }
        lower = text.lower()
        for field, labels in section_groups.items():
            if not any(label in lower for label in labels):
                issues.append(
                    ResumeValidationIssue(
                        code=f"missing_{field}",
                        severity=ResumeIssueSeverity.INFO,
                        field=field,
                        message=f"未识别到{labels[0]}版块",
                    )
                )
        current_year = datetime.now(timezone.utc).year
        future_years = sorted(
            {
                int(year)
                for year in re.findall(r"\b(?:19|20)\d{2}\b", text)
                if int(year) > current_year
            }
        )
        if future_years:
            issues.append(
                ResumeValidationIssue(
                    code="future_date",
                    severity=ResumeIssueSeverity.WARNING,
                    field="timeline",
                    message=f"发现未来年份：{', '.join(map(str, future_years))}，请确认时间线",
                )
            )
        overlaps = ResumeProcessor._find_employment_overlaps(text)
        if overlaps:
            examples = "；".join(f"{left} ↔ {right}" for left, right in overlaps[:3])
            issues.append(
                ResumeValidationIssue(
                    code="overlapping_employment",
                    severity=ResumeIssueSeverity.WARNING,
                    field="timeline",
                    message=f"发现可能重叠的任职时间，请确认是否为兼职、顾问或并行任职：{examples}",
                )
            )
        if ResumeProcessor._contains_instruction_like_text(text):
            issues.append(
                ResumeValidationIssue(
                    code="instruction_like_text",
                    severity=ResumeIssueSeverity.WARNING,
                    field="document",
                    message=(
                        "发现疑似面向 AI 的指令文本；系统会将其仅作为不可信简历内容，"
                        "不会执行其中指令，请核对是否属于简历正文"
                    ),
                )
            )
        return issues

    @staticmethod
    def _contains_instruction_like_text(text: str) -> bool:
        """Detect common prompt-injection wording without interpreting the document."""
        patterns = (
            r"(?im)^\s*(?:system|assistant|developer)\s*[:：]",
            r"(?i)ignore\s+(?:all\s+)?(?:previous|prior|above)\s+instructions?",
            r"(?i)(?:mark|set|rate|score).{0,40}(?:confirmed|100|full\s+marks)",
            r"(?:忽略|无视).{0,12}(?:此前|之前|以上|所有).{0,8}(?:指令|规则|要求)",
            r"(?:将|把).{0,30}(?:标记为已确认|评分为?\s*100|满分)",
        )
        return any(re.search(pattern, text) for pattern in patterns)

    @staticmethod
    def _find_employment_overlaps(text: str) -> list[tuple[str, str]]:
        """Return plausible overlaps between full-time employment date ranges.

        The result is a review prompt, not a fraud verdict: concurrent work can
        be legitimate. Education, internships, consulting, and explicit
        part-time/advisory entries are excluded to reduce false positives.
        """
        intervals = ResumeProcessor._employment_intervals(text)
        intervals.sort(key=lambda item: (item.start_month, item.end_month))
        overlaps: list[tuple[str, str]] = []
        furthest_ending: _EmploymentInterval | None = None
        for current in intervals:
            if furthest_ending is not None and furthest_ending.end_month >= current.start_month:
                # Ignore a boundary-month handoff; date-only resumes commonly
                # use the same month for one role ending and the next starting.
                overlap_months = (
                    min(furthest_ending.end_month, current.end_month) - current.start_month
                )
                same_boundary_year = (
                    furthest_ending.end_month // 12 == current.start_month // 12
                    and (
                        not furthest_ending.end_is_precise or not current.start_is_precise
                    )
                )
                if overlap_months >= 1 and not same_boundary_year:
                    overlaps.append((furthest_ending.statement, current.statement))
            if furthest_ending is None or current.end_month > furthest_ending.end_month:
                furthest_ending = current
        return overlaps

    @staticmethod
    def _employment_intervals(text: str) -> list[_EmploymentInterval]:
        month_names = {
            name: index
            for index, names in enumerate(
                (
                    ("jan", "january"),
                    ("feb", "february"),
                    ("mar", "march"),
                    ("apr", "april"),
                    ("may",),
                    ("jun", "june"),
                    ("jul", "july"),
                    ("aug", "august"),
                    ("sep", "sept", "september"),
                    ("oct", "october"),
                    ("nov", "november"),
                    ("dec", "december"),
                ),
                start=1,
            )
            for name in names
        }
        month_pattern = "|".join(sorted(month_names, key=len, reverse=True))
        named_range = re.compile(
            rf"\b({month_pattern})\.?\s+((?:19|20)\d{{2}})\s*"
            rf"(?:-|–|—|to|至)\s*(present|current|now|至今|({month_pattern})\.?\s+((?:19|20)\d{{2}}))",
            re.IGNORECASE,
        )
        numeric_range = re.compile(
            r"\b((?:19|20)\d{2})(?:[-/.](\d{1,2}))?\s*"
            r"(?:-|–|—|to|至)\s*(present|current|now|至今|((?:19|20)\d{2})(?:[-/.](\d{1,2}))?)",
            re.IGNORECASE,
        )
        excluded = re.compile(
            r"education|university|college|school|degree|bachelor|master|ph\.?d|"
            r"教育|大学|学院|学校|本科|硕士|博士|intern|实习|part[ -]?time|兼职|"
            r"consult(?:ant|ing)?|顾问|advisor|adviser",
            re.IGNORECASE,
        )
        experience_heading = re.compile(
            r"^(?:experience|employment|work experience|professional experience|"
            r"工作经历|工作经验|职业经历)$",
            re.IGNORECASE,
        )
        other_heading = re.compile(
            r"^(?:summary|profile|education|skills?|capabilities|competencies|expertise|"
            r"projects?|research|leadership|awards?|certifications?|languages?|"
            r"个人简介|教育经历|教育背景|技能|专业能力|核心能力|项目经历|项目经验|"
            r"研究经历|领导力|奖项|证书|语言)$",
            re.IGNORECASE,
        )
        intervals: list[_EmploymentInterval] = []
        merged_lines = ResumeProcessor._merge_vertical_date_ranges(text.splitlines())
        has_experience_section = any(
            experience_heading.fullmatch(line.strip().rstrip(":：")) for line in merged_lines
        )
        in_experience_section = not has_experience_section
        for raw_line in merged_lines:
            line = raw_line.strip()
            heading = line.rstrip(":：")
            if experience_heading.fullmatch(heading):
                in_experience_section = True
                continue
            if other_heading.fullmatch(heading):
                in_experience_section = False
                continue
            if not in_experience_section:
                continue
            if not line or excluded.search(line):
                continue
            match = named_range.search(line)
            if match:
                start = int(match.group(2)) * 12 + month_names[match.group(1).lower()] - 1
                if match.group(3).lower() in {"present", "current", "now", "至今"}:
                    now = datetime.now(timezone.utc)
                    end = now.year * 12 + now.month - 1
                else:
                    end = int(match.group(5)) * 12 + month_names[match.group(4).lower()] - 1
                intervals.append(_EmploymentInterval(start, end, line[:180], True, True))
                continue
            match = numeric_range.search(line)
            if not match:
                continue
            start = int(match.group(1)) * 12 + int(match.group(2) or 1) - 1
            if match.group(3).lower() in {"present", "current", "now", "至今"}:
                now = datetime.now(timezone.utc)
                end = now.year * 12 + now.month - 1
                end_is_precise = True
            else:
                end = int(match.group(4)) * 12 + int(match.group(5) or 12) - 1
                end_is_precise = bool(match.group(5))
            intervals.append(
                _EmploymentInterval(start, end, line[:180], bool(match.group(2)), end_is_precise)
            )
        return intervals

    @staticmethod
    def _find_claims(text: str) -> list[ResumeClaim]:
        claims: list[ResumeClaim] = []
        patterns = (
            (
                "education",
                re.compile(
                    r"大学|学院|university|college|bachelor|master|博士|硕士|本科", re.IGNORECASE
                ),
                "supporting_document",
            ),
            (
                "employment",
                re.compile(
                    r"有限公司|集团|inc\.?|corp\.?|company|任职|就职|"
                    r"(?:19|20)\d{2}[-/.]\d{1,2}\s+to\s+[A-Za-z]",
                    re.IGNORECASE,
                ),
                "public_source_or_reference",
            ),
            (
                "certification",
                re.compile(r"证书|认证|certified|certification|certificate", re.IGNORECASE),
                "supporting_document",
            ),
            (
                "achievement",
                re.compile(r"\d+(?:\.\d+)?\s*(?:%|倍|x\b|万|亿|million|k\b)", re.IGNORECASE),
                "candidate_confirmation",
            ),
        )
        seen: set[tuple[str, str]] = set()
        for line in ResumeProcessor._merge_vertical_date_ranges(text.splitlines()):
            statement = line.strip()
            if not 8 <= len(statement) <= 300:
                continue
            degree_line = bool(
                re.search(
                    r"\b(?:bachelor|master|ph\.?d\.?|doctorate|degree)\b|"
                    r"博士|硕士|本科|学位",
                    statement,
                    re.IGNORECASE,
                )
            )
            for category, pattern, method in patterns:
                # A dated degree line often also matches the generic English
                # employment timeline pattern ("YYYY-MM to Institution"). It
                # is one education claim, not simultaneous employment.
                if category == "employment" and degree_line:
                    continue
                if category == "employment" and not re.search(
                    r"\b(?:19|20)\d{2}\b|\d{4}[-/.]\d{1,2}", statement
                ):
                    continue
                key = (category, statement)
                if pattern.search(statement) and key not in seen:
                    claims.append(
                        ResumeClaim(
                            category=category,
                            statement=statement,
                            original_statement=statement,
                            verification_method=method,
                        )
                    )
                    seen.add(key)
                    if len(claims) == 40:
                        return claims
        return claims

    @staticmethod
    def _merge_vertical_date_ranges(lines: list[str]) -> list[str]:
        """Fold a lone end-date line into the preceding 'YYYY-MM to Company' line.

        Some PDFs lay out a date range vertically: the start date + employer on
        one line, the end date alone on the next. Without merging, the end date
        becomes an orphan line and timeline claims lose their period.
        """
        merged: list[str] = []
        for line in lines:
            statement = line.strip()
            prev = merged[-1] if merged else ""
            # Only fold when the previous line is a date-range start
            # ("YYYY-MM to Company"), not any line that happens to contain "to".
            is_range_start = bool(
                re.search(
                    r"\b(?:19|20)\d{2}[-/.]\d{1,2}\s+(?:to|至)\s+[A-Za-z一-鿿]",
                    prev,
                )
            )
            if (
                statement
                and merged
                and re.fullmatch(r"(?:19|20)\d{2}[-/.]\d{1,2}", statement)
                and is_range_start
                and not re.search(r"(?:19|20)\d{2}[-/.]\d{1,2}\s*to\s*\S+\s*(?:19|20)\d{2}[-/.]\d{1,2}", prev)
            ):
                merged[-1] = f"{merged[-1]} {statement}"
                continue
            merged.append(line)
        return merged
