import json
from io import BytesIO
from uuid import uuid4

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from interview_os.api.app import create_app
from interview_os.database.storage import Storage
from interview_os.models.local_llm import LocalLLMClient
from interview_os.services.interview_service import InterviewService
from interview_os.services.settings_service import LocalSettingsStore
from interview_os.tools.asr import ASRClient
from interview_os.tools.web_search import SearchResult


def _register_live_capture(client: TestClient, session_id: str) -> str:
    settings = client.get("/api/settings").json()
    live = client.get(f"/api/live-interviews/{session_id}").json()["state"][
        "live_interview"
    ]
    recording_id = str(uuid4())
    response = client.post(
        f"/api/live-interviews/{session_id}/audio/captures/{recording_id}",
        json={
            "expected_capture_epoch": live["capture_epoch"],
            "expected_settings_revision": settings["audio_settings_revision"],
            "expected_settings_etag": settings["audio_settings_etag"],
        },
    )
    assert response.status_code == 200
    return recording_id


class MockLLM:
    async def chat(self, messages, **kwargs):
        return '{"name":"Ada","skills":["Python"],"strengths":["Systems"]}'

    async def embed(self, text):
        return []


class WorkflowLLM:
    async def chat(self, messages, **kwargs):
        prompt = messages[-1]["content"].lower()
        if "structured candidate profile" in prompt:
            return '{"name":"Ada","skills":["Python"]}'
        if "job description" in prompt:
            return '{"title":"Engineer","competencies":["System Design"]}'
        if "company dna" in prompt:
            return '{"name":"Example","dna":"Engineering culture"}'
        if "fuse three inputs" in prompt:
            return '{"summary":"Show impact","key_risks":["Scope"]}'
        if "mock interview plan" in prompt:
            return '{"questions":[{"question":"Design it","competency":"System Design"}]}'
        if "你是面试问题分析器" in messages[-1]["content"]:
            return '{"answer_type":"behavioral_example","answer_type_label":"行为经历题","assessment_goal":"验证复盘能力是否有真实证据","competency":"复盘能力","answer_boundary":["提供真实案例"],"common_mistakes":["只讲团队"],"transfer_principle":"复用事实，切换取舍、修正或协作重点。","related_questions":["讲一次关键取舍。","讲一次失败后的调整。","同事会如何描述你的贡献？"],"likely_follow_ups":["你个人做了什么？","结果如何验证？"]}'
        if "analyze this interview answer" in prompt:
            return '{"content":0.8,"technical_depth":0.8,"structure":0.8,"impact":0.8,"feedback":[],"improved_answer":"Better","observed_signals":["Clear design"],"missing_signals":[]}'
        if "based on the following evidence" in prompt:
            return '{"competencies":[{"competency":"System Design","score":0.8,"confidence":0.8,"supporting_evidence":["Clear design"],"gaps":[]}],"overall_score":0.8,"recommendation":"hire","summary":"Meets the bar","risks":[]}'
        if "generate evidence-based feedback" in prompt:
            return '{"overall":"Meets the bar","strengths":["Clear design"],"improvements":[],"action_plan":["Continue practice"],"interviewer_notes":[],"recommendation_reasoning":"Evidence supports hire."}'
        if "return exactly one json object matching the questionsuggestion schema" in prompt:
            return '{"suggested_question":"你如何验证这个架构权衡？","question_type":"follow_up","competency":"System Design","rationale":"需要补充验证方法","evidence_gap":"量化验证","expected_signals":["指标","压测"],"confidence":0.8,"alternatives":["失败时如何回滚？"]}'
        if "参考答案提示框架" in prompt:
            return ('{"frameworks":['
                    '{"question_index":0,"answer_framework":"用 STAR 讲 2024 事故复盘"},'
                    '{"question_index":1,"answer_framework":"用 SLO 故事量化平台健康"},'
                    '{"question_index":2,"answer_framework":"讲清取舍并给量化结果"},'
                    '{"question_index":3,"answer_framework":"引用 ZUORA 账单平台扩展经历"},'
                    '{"question_index":4,"answer_framework":"突出团队从 4 人扩展到 15 人的领导力"}]}')
        raise AssertionError(prompt)

    async def embed(self, text):
        return []


class FollowupWorkflowLLM(WorkflowLLM):
    async def chat(self, messages, **kwargs):
        prompt = messages[-1]["content"].lower()
        if "mock interview plan" in prompt:
            return '{"questions":[{"question":"Design it","competency":"System Design","follow_ups":["What evidence supports that trade-off?"]}]}'
        if "analyze this interview answer" in prompt:
            return '{"content":0.6,"technical_depth":0.6,"structure":0.6,"impact":0.5,"feedback":["Add evidence"],"improved_answer":"Better","observed_signals":["Design thinking"],"missing_signals":["Measured impact"]}'
        return await super().chat(messages, **kwargs)


class BlueprintFallbackQuestionLLM(WorkflowLLM):
    async def chat(self, messages, **kwargs):
        prompt = messages[-1]["content"].lower()
        if "interview blueprint" in prompt:
            return (
                '{"position":"Engineer","rounds":[{"name":"技术面","goal":"验证核心能力",'
                '"evaluation_criteria":["证据完整"],"questions":['
                '{"question":"请讲一次系统设计权衡。","competency":"System Design",'
                '"rationale":"核心能力","strong_signals":["权衡","指标"],"follow_ups":[]},'
                '{"question":"请讲一次事故复盘。","competency":"Incident Review",'
                '"rationale":"稳定性能力","strong_signals":["复盘","改进"],"follow_ups":[]}]}]}'
            )
        if (
            "return exactly one json object matching the questionsuggestion schema" in prompt
            or "repair the previous response into valid json only" in prompt
        ):
            return "not json"
        return await super().chat(messages, **kwargs)


class FactSearchProvider:
    async def search(self, query, limit=5, *, search_depth="basic"):
        return [
            SearchResult(
                title="Example AI platform",
                url="https://example.com/about",
                snippet="Example develops AI systems for enterprise interview workflows.",
                source="test",
                source_quality="official",
                is_official=True,
            )
        ][:limit]


def make_resume_docx() -> bytes:
    document = Document()
    document.add_paragraph("Ada ada@example.com 13800138000")
    document.add_paragraph("教育经历 Example University 本科")
    document.add_paragraph("工作经历 Example有限公司 将性能提升 30%")
    document.add_paragraph("技能 Python FastAPI")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_resume_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Resume", styles["Title"]),
        Paragraph("SMIC High School 08/2020 - 06/2024", styles["BodyText"]),
        Paragraph("ZUORA Senior Recruiting Manager 2018-07 - 2022-12", styles["BodyText"]),
    ]
    document.build(story)
    return buffer.getvalue()


def make_scanned_resume_pdf() -> bytes:
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    image = Image.new("RGB", (1000, 1400), "white")
    ImageDraw.Draw(image).text((60, 80), "SCANNED RESUME", fill="black")
    image_buffer = BytesIO()
    image.save(image_buffer, format="PNG")
    pdf_buffer = BytesIO()
    pdf = canvas.Canvas(pdf_buffer, pagesize=A4)
    width, height = A4
    pdf.drawImage(ImageReader(BytesIO(image_buffer.getvalue())), 0, 0, width, height)
    pdf.showPage()
    pdf.save()
    return pdf_buffer.getvalue()


def test_session_resume_analysis_flow(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'api.db'}")
    app = create_app(storage=storage, llm_client=MockLLM(), configure_llm=False)

    with TestClient(app) as client:
        home = client.get("/")
        assert home.status_code == 200
        assert "面试智能工作台" in home.text
        script = client.get("/static/app.js")
        assert script.status_code == 200
        assert client.get("/app.js").status_code == 200
        assert client.get("/styles.css").status_code == 200
        assert client.get("/modules/runtime.js").status_code == 200
        candidate_view = client.get("/static/modules/candidate-view.js")
        assert candidate_view.status_code == 200
        assert "本人确认 · 未外部核验" in candidate_view.text
        assert "恢复待核验" in candidate_view.text
        assert "restoreAuthenticatedSession" in script.text
        assert "displayedQuestionText = justAnswered ? displayResponse.question" in script.text
        assert "mock-retry-main" in script.text
        assert "displayedAsFollowUp = justAnswered ? displayResponse.is_follow_up" in script.text
        assert "resetMockAudioExperience" in script.text
        assert "recordingSessionId!==state.sessionId" in script.text
        assert "response_id=${encodeURIComponent(responseId)}" in script.text
        assert "requestedSpeechKey!==mockCurrentSpeechKey" in script.text
        info = client.get("/", headers={"Accept": "application/json"})
        assert info.json()["name"] == "InterviewOS"

        created = client.post(
            "/api/interviews/sessions",
            json={"candidate_name": "Initial", "job_title": "Engineer"},
        )
        assert created.status_code == 200
        session_id = created.json()["id"]

        analyzed = client.post(
            "/api/analysis/resume",
            json={"session_id": session_id, "text": "Experienced Python developer"},
        )
        assert analyzed.status_code == 200
        assert analyzed.json()["state"]["candidate"]["name"] == "Ada"

        fetched = client.get(f"/api/interviews/sessions/{session_id}")
        assert fetched.status_code == 200
        assert fetched.json()["state"]["candidate"]["skills"] == ["Python"]

        debug = client.get("/api/debug/events")
        assert debug.status_code == 200
        actions = [event["action"] for event in debug.json()["events"]]
        assert "session_created" in actions
        assert "agent_completed" in actions


def test_missing_session_returns_404(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'missing.db'}")
    app = create_app(storage=storage, llm_client=MockLLM(), configure_llm=False)
    with TestClient(app) as client:
        response = client.get("/api/interviews/sessions/missing")
    assert response.status_code == 404


def test_resume_upload_and_human_confirmation_flow(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-upload.db'}")
    app = create_app(storage=storage, llm_client=MockLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        uploaded = client.post(
            f"/api/resumes/{session_id}/upload",
            files={
                "file": (
                    "ada.docx",
                    make_resume_docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert uploaded.status_code == 200
        state = uploaded.json()["state"]
        assert "Ada" in state["candidate"]["raw_resume_text"]
        claim = state["resume_review"]["claims"][0]

        confirmed = client.patch(
            f"/api/resumes/{session_id}/claims/{claim['id']}",
            json={"status": "confirmed", "note": "候选人已确认"},
        )
        restored = client.patch(
            f"/api/resumes/{session_id}/claims/{claim['id']}",
            json={"status": "unverified"},
        )
        invalid = client.post(
            f"/api/resumes/{session_id}/upload",
            files={"file": ("old.doc", b"legacy", "application/msword")},
        )

    assert confirmed.status_code == 200
    assert confirmed.json()["state"]["resume_review"]["claims"][0]["status"] == "confirmed"
    assert restored.status_code == 200
    assert restored.json()["state"]["resume_review"]["claims"][0]["status"] == "unverified"
    assert invalid.status_code == 422


def test_settings_api_configures_tavily_without_exposing_key(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'settings.db'}")
    llm = LocalLLMClient(base_url="http://localhost:11434/v1", api_key="local", model="test")
    app = create_app(
        storage=storage,
        llm_client=llm,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        updated = client.put(
            "/api/settings",
            json={
                "search": {"provider": "tavily", "tavily_api_key": "tvly-secret"},
                "tts": {
                    "base_url": "http://tts.test/v1",
                    "model": "qwen3-tts",
                    "api_key": "tts-secret",
                },
            },
        )
        assert updated.status_code == 200
        body = updated.json()
        assert body["search"]["selected"] == "tavily"
        assert body["search"]["configured"]["tavily"] is True
    assert "tvly-secret" not in updated.text
    assert "tts-secret" not in updated.text
    assert (tmp_path / "settings.json").stat().st_mode & 0o777 == 0o600


def test_settings_failure_rolls_back_earlier_provider_mutations(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'settings-rollback.db'}")
    llm = LocalLLMClient(base_url="http://llm.test/v1", api_key="local", model="main")
    app = create_app(
        storage=storage,
        llm_client=llm,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        failed = client.put(
            "/api/settings",
            json={
                "search": {"provider": "tavily", "tavily_api_key": "temporary-secret"},
                "asr": {"base_url": "ftp://invalid-asr"},
                "persist": False,
            },
        )
        assert failed.status_code == 400
        current = client.get("/api/settings").json()
        assert current["search"]["selected"] == "none"
        assert current["search"]["configured"]["tavily"] is False
        assert current["asr"]["base_url"] == "http://192.168.1.97:8007"


@pytest.mark.parametrize(
    "payload",
    [
        {"llm": {"base_url": "file:///tmp/model"}},
        {"asr": {"base_url": ""}},
        {"asr": {"base_url": "http://"}},
        {"live_audio": {"base_url": "https:///missing-host"}},
        {"resume_llm": {"base_url": "file://localhost/model"}},
        {"tts": {"base_url": "http://tts.test:bad"}},
        {"asr": {"base_url": "http://\ud800"}},
        {"live_audio": {"base_url": "http://exa\ud800mple.test/v1"}},
        {
            "search": {
                "provider": "searxng",
                "searxng_base_url": "file:///tmp/search",
            }
        },
    ],
)
def test_settings_put_reuses_startup_http_endpoint_policy(tmp_path, payload):
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'invalid-url.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )

    with TestClient(app) as client:
        rejected = client.put("/api/settings", json={**payload, "persist": False})

    assert rejected.status_code == 400
    assert "http:// or https://" in rejected.json()["detail"]


@pytest.mark.parametrize(
    "payload",
    [
        {"search": {"provider": "tavily", "tavily_api_key": "bad\nkey"}},
        {"search": {"provider": "brave", "brave_api_key": "密钥"}},
        {"llm": {"api_key": " leading"}},
        {"asr": {"api_key": "bad key"}},
        {"live_audio": {"api_key": "bad\tkey"}},
        {"resume_llm": {"api_key": "trailing "}},
        {"tts": {"api_key": "\ud800"}},
    ],
)
def test_settings_reject_header_unsafe_credentials_before_runtime_mutation(
    tmp_path, payload
):
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'invalid-key.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    before = app.state.asr_client.secret_snapshot()

    with TestClient(app) as client:
        rejected = client.put(
            "/api/settings", json={**payload, "persist": False}
        )

    assert rejected.status_code == 400
    assert "API key" in rejected.json()["detail"]
    assert app.state.asr_client.secret_snapshot() == before


def test_settings_persistence_failure_restores_runtime_configuration(tmp_path):
    class FailingSettingsStore(LocalSettingsStore):
        def save(self, payload):
            raise OSError("private filesystem detail")

    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'settings-save-rollback.db'}")
    llm = LocalLLMClient(base_url="http://llm.test/v1", api_key="local", model="main")
    app = create_app(
        storage=storage,
        llm_client=llm,
        configure_llm=False,
        settings_store=FailingSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        failed = client.put(
            "/api/settings",
            json={"llm": {"model": "must-not-stick"}, "persist": True},
        )
        assert failed.status_code == 500
        assert failed.json()["detail"] == "设置保存失败，运行时配置未更改"
        assert client.app.state.llm_client.model == "main"


def test_resume_client_rollback_publishes_original_before_failed_cleanup(
    tmp_path, monkeypatch
):
    class FailingSettingsStore(LocalSettingsStore):
        def save(self, payload):
            raise OSError("simulated persistence failure")

    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-rollback.db'}")
    primary = LocalLLMClient(base_url="http://main.test/v1", model="main")
    original_close = LocalLLMClient.close

    async def fail_split_close(self):
        if self is not primary:
            raise RuntimeError("simulated cleanup failure")
        await original_close(self)

    monkeypatch.setattr(LocalLLMClient, "close", fail_split_close)
    app = create_app(
        storage=storage,
        llm_client=primary,
        configure_llm=False,
        settings_store=FailingSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        failed = client.put(
            "/api/settings",
            json={
                "resume_llm": {
                    "base_url": "http://resume.test/v1",
                    "model": "resume",
                },
                "persist": True,
            },
        )

        assert failed.status_code == 500
        assert client.app.state.resume_llm_client is primary
        assert client.app.state.interview_service.resume_llm_client is primary


def test_failed_settings_update_is_rolled_back_before_event_loop_observers(
    tmp_path,
):
    from threading import Event

    observer_ran = Event()

    def observe_url():
        observed_urls.append(llm.base_url)
        observer_ran.set()

    class ObserverFailingStore(LocalSettingsStore):
        fail = False

        def save(self, payload):
            if not self.fail:
                return super().save(payload)
            import asyncio

            loop = asyncio.get_running_loop()
            loop.call_soon(observe_url)
            raise OSError("simulated persistence failure")

    observed_urls = []
    llm = LocalLLMClient(base_url="http://committed.test/v1", model="main")
    store = ObserverFailingStore(tmp_path / "settings.json")
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'settings-isolation.db'}"),
        llm_client=llm,
        configure_llm=False,
        settings_store=store,
    )
    with TestClient(app) as client:
        store.fail = True
        failed = client.put(
            "/api/settings",
            json={"llm": {"base_url": "http://uncommitted.test/v1"}},
        )

        assert failed.status_code == 500
        assert observer_ran.wait(timeout=1)
        assert observed_urls == ["http://committed.test/v1"]
        assert llm.base_url == "http://committed.test/v1"


def test_startup_read_error_never_overwrites_existing_secret_settings(tmp_path):
    settings_path = tmp_path / "settings.json"
    original = '{"search":{"provider":"tavily","tavily_api_key":"keep-me"}}'
    settings_path.write_text(original, encoding="utf-8")

    class ReadErrorStore(LocalSettingsStore):
        def load_with_status(self):
            return {}, "error"

    create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'read-error.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=ReadErrorStore(settings_path),
    )

    assert settings_path.read_text(encoding="utf-8") == original


@pytest.mark.parametrize(
    "original",
    [
        '{"llm":{"input_cost_per_million":Infinity}}',
        '{"llm":{"input_cost_per_million":1e309}}',
    ],
)
def test_startup_preserves_non_finite_legacy_settings_for_manual_recovery(
    tmp_path, original
):
    settings_path = tmp_path / "settings.json"
    settings_path.write_text(original, encoding="utf-8")

    create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'invalid-cost.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )

    assert settings_path.read_text(encoding="utf-8") == original


@pytest.mark.parametrize(
    "payload",
    [
        {"asr": None},
        {"asr": {"base_url": ""}},
        {"llm": []},
        {"tts": "bad"},
        {"live_audio": {"base_url": "not-a-url"}},
        {"resume_llm": {"unknown": "field"}},
    ],
)
def test_startup_ignores_malformed_settings_sections_without_overwriting(
    tmp_path, payload
):
    settings_path = tmp_path / "settings.json"
    original = json.dumps(payload)
    settings_path.write_text(original, encoding="utf-8")

    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'invalid-section.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )
    with TestClient(app) as client:
        response = client.get("/api/settings")

    assert response.status_code == 200
    assert settings_path.read_text(encoding="utf-8") == original


def test_startup_migrates_known_obsolete_resume_llm_fields(tmp_path):
    settings_path = tmp_path / "settings.json"
    settings_path.write_text(
        json.dumps(
            {
                "resume_llm": {
                    "base_url": "http://resume.test/v1",
                    "model": "omni",
                    "embedding_model": "unused",
                    "input_cost_per_million": 0,
                    "output_cost_per_million": 0,
                }
            }
        ),
        encoding="utf-8",
    )
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'legacy-resume.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )

    with TestClient(app) as client:
        assert client.get("/api/settings").status_code == 200

    persisted = json.loads(settings_path.read_text(encoding="utf-8"))
    assert persisted["resume_llm"]["base_url"] == "http://resume.test/v1"
    assert "embedding_model" not in persisted["resume_llm"]
    assert "input_cost_per_million" not in persisted["resume_llm"]
    assert "output_cost_per_million" not in persisted["resume_llm"]


def test_settings_reject_non_finite_or_unbounded_costs_without_committing(tmp_path):
    settings_path = tmp_path / "settings.json"
    llm = LocalLLMClient(base_url="http://llm.test/v1", model="main")
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'cost-bounds.db'}"),
        llm_client=llm,
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )
    with TestClient(app) as client:
        original = settings_path.read_text(encoding="utf-8")
        infinite = client.put(
            "/api/settings",
            content='{"llm":{"input_cost_per_million":1e309}}',
            headers={"Content-Type": "application/json"},
        )
        excessive = client.put(
            "/api/settings",
            json={"search": {"provider": "none", "search_request_cost_usd": 1e308}},
        )

        assert infinite.status_code == 422
        assert excessive.status_code == 422
        assert llm.input_cost_per_million == 0
        assert settings_path.read_text(encoding="utf-8") == original


def test_audio_settings_generation_and_pending_capture_barrier(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'audio-settings-cas.db'}")
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        initial = client.get("/api/settings").json()
        assert len(initial["audio_settings_etag"]) == 32
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )

        changed = client.put(
            "/api/settings",
            json={"asr": {"model": "new-asr"}, "persist": False},
        )
        assert changed.status_code == 200
        current = changed.json()
        assert current["audio_settings_revision"] == initial["audio_settings_revision"] + 1
        assert current["audio_settings_etag"] != initial["audio_settings_etag"]

        recording_id = uuid4()
        stale = client.post(
            f"/api/live-interviews/{session_id}/audio/captures/{recording_id}",
            json={
                "expected_capture_epoch": 1,
                "expected_settings_revision": initial["audio_settings_revision"],
                "expected_settings_etag": initial["audio_settings_etag"],
            },
        )
        registered = client.post(
            f"/api/live-interviews/{session_id}/audio/captures/{recording_id}",
            json={
                "expected_capture_epoch": 1,
                "expected_settings_revision": current["audio_settings_revision"],
                "expected_settings_etag": current["audio_settings_etag"],
            },
        )
        no_op_audio = client.put(
            "/api/settings",
            json={
                "search": {"provider": "none", "search_request_cost_usd": 0.25},
                "asr": {
                    "base_url": current["asr"]["base_url"],
                    "model": current["asr"]["model"],
                    "transcription_path": current["asr"]["transcription_path"],
                    "timeout_seconds": current["asr"]["timeout_seconds"],
                },
                "live_audio": {
                    "mode": current["live_audio"]["mode"],
                    "name": current["live_audio"]["name"],
                    "base_url": current["live_audio"]["base_url"],
                    "model": current["live_audio"]["model"],
                },
                "persist": False,
            },
        )
        blocked = client.put(
            "/api/settings",
            json={"asr": {"model": "blocked-asr"}, "persist": False},
        )
        cancelled = client.delete(
            f"/api/live-interviews/{session_id}/audio/captures/{recording_id}"
        )
        unblocked = client.put(
            "/api/settings",
            json={"asr": {"model": "unblocked-asr"}, "persist": False},
        )

        assert stale.status_code == 409
        assert registered.status_code == 200
        assert no_op_audio.status_code == 200
        assert no_op_audio.json()["audio_settings_revision"] == current[
            "audio_settings_revision"
        ]
        assert no_op_audio.json()["audio_settings_etag"] == current[
            "audio_settings_etag"
        ]
        assert no_op_audio.json()["search"]["search_request_cost_usd"] == 0.25
        assert blocked.status_code == 409
        assert cancelled.status_code == 200
        assert unblocked.status_code == 200


def test_invalid_search_secret_does_not_poison_runtime_settings(tmp_path):
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'search-atomic.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        accepted = client.put(
            "/api/settings",
            json={
                "search": {
                    "provider": "tavily",
                    "tavily_api_key": "working-key",
                },
                "persist": False,
            },
        )
        before = app.state.search_manager.secret_snapshot()
        rejected = client.put(
            "/api/settings",
            json={
                "search": {
                    "provider": "tavily",
                    "tavily_api_key": "\ud800",
                },
                "persist": False,
            },
        )
        after = app.state.search_manager.secret_snapshot()
        recovered = client.put(
            "/api/settings",
            json={
                "search": {
                    "provider": "tavily",
                    "tavily_api_key": "replacement-key",
                },
                "persist": False,
            },
        )

    assert accepted.status_code == 200
    assert rejected.status_code == 400
    assert after == before
    assert recovered.status_code == 200


def test_provider_credentials_can_be_explicitly_cleared_and_removed_from_disk(
    tmp_path,
):
    settings_path = tmp_path / "settings.json"
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'clear-secrets.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )
    with TestClient(app) as client:
        configured = client.put(
            "/api/settings",
            json={
                "search": {
                    "provider": "tavily",
                    "tavily_api_key": "tavily-old-secret",
                },
                "asr": {"api_key": "asr-old-secret"},
                "live_audio": {"api_key": "omni-old-secret"},
                "tts": {"api_key": "tts-old-secret"},
            },
        )
        cleared = client.put(
            "/api/settings",
            json={
                "search": {
                    "provider": "none",
                    "tavily_api_key": "",
                    "brave_api_key": "",
                },
                "asr": {"api_key": ""},
                "live_audio": {"api_key": ""},
                "tts": {"api_key": ""},
            },
        )

    assert configured.status_code == 200
    assert cleared.status_code == 200
    assert cleared.json()["search"]["configured"]["tavily"] is False
    assert app.state.search_manager.secret_snapshot()["tavily_api_key"] == ""
    assert app.state.asr_client.secret_snapshot()["api_key"] == ""
    assert app.state.omni_client.secret_snapshot()["api_key"] == ""
    assert app.state.tts_client.secret_snapshot()["api_key"] == ""
    persisted = settings_path.read_text(encoding="utf-8")
    for secret in (
        "tavily-old-secret",
        "asr-old-secret",
        "omni-old-secret",
        "tts-old-secret",
    ):
        assert secret not in persisted


def test_persistent_audio_generation_survives_backend_restart(tmp_path):
    database_url = f"sqlite+aiosqlite:///{tmp_path / 'audio-settings-restart.db'}"
    settings_path = tmp_path / "settings.json"
    recording_id = uuid4()

    first_app = create_app(
        storage=Storage(database_url),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )
    with TestClient(first_app) as client:
        changed = client.put(
            "/api/settings",
            json={"asr": {"model": "persistent-asr"}, "persist": True},
        ).json()
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        started = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        ).json()
        capture_epoch = started["state"]["live_interview"]["capture_epoch"]
        registered = client.post(
            f"/api/live-interviews/{session_id}/audio/captures/{recording_id}",
            json={
                "expected_capture_epoch": capture_epoch,
                "expected_settings_revision": changed["audio_settings_revision"],
                "expected_settings_etag": changed["audio_settings_etag"],
            },
        )
        assert registered.status_code == 200

    second_app = create_app(
        storage=Storage(database_url),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )
    with TestClient(second_app) as client:
        restored = client.get("/api/settings").json()
        assert restored["audio_settings_etag"] == changed["audio_settings_etag"]
        assert restored["audio_settings_revision"] == changed["audio_settings_revision"]
        renewed = client.post(
            f"/api/live-interviews/{session_id}/audio/captures/{recording_id}",
            json={
                "expected_capture_epoch": capture_epoch,
                "expected_settings_revision": restored["audio_settings_revision"],
                "expected_settings_etag": restored["audio_settings_etag"],
            },
        )
        assert renewed.status_code == 200


def test_resume_model_settings_split_default_alias_from_primary_llm(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-model-split.db'}")
    primary = LocalLLMClient(
        base_url="http://main.test/v1", api_key="main-key", model="main-model"
    )
    app = create_app(
        storage=storage,
        llm_client=primary,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        assert client.app.state.resume_llm_client is primary
        updated = client.put(
            "/api/settings",
            json={
                "resume_llm": {
                    "base_url": "http://resume.test/v1",
                    "model": "resume-model",
                },
                "persist": False,
            },
        )
        assert updated.status_code == 200
        dedicated = client.app.state.resume_llm_client
        assert dedicated is not primary
        assert dedicated.model == "resume-model"
        assert dedicated.base_url == "http://resume.test/v1"
        assert primary.model == "main-model"
        assert primary.base_url == "http://main.test/v1"
        assert client.app.state.interview_service.resume_llm_client is dedicated


def test_resume_model_split_preserves_injected_transport(tmp_path):
    import httpx

    requests = []

    def handler(request):
        requests.append(str(request.url))
        return httpx.Response(200, json={"data": [{"id": "resume-model"}]})

    primary = LocalLLMClient(
        base_url="http://main.test/v1",
        model="main",
        transport=httpx.MockTransport(handler),
    )
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-transport.db'}"),
        llm_client=primary,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        updated = client.put(
            "/api/settings",
            json={
                "resume_llm": {
                    "base_url": "http://resume.test/v1",
                    "model": "resume",
                },
                "persist": False,
            },
        )
        assert updated.status_code == 200
        dedicated = client.app.state.resume_llm_client
        result = client.portal.call(dedicated.probe)

        assert result["models"] == ["resume-model"]
        assert requests == ["http://resume.test/v1/models"]


def test_persisted_resume_model_preserves_injected_transport_after_restart(tmp_path):
    import httpx

    settings_path = tmp_path / "settings.json"
    first_primary = LocalLLMClient(
        base_url="http://main.test/v1",
        transport=httpx.MockTransport(
            lambda request: httpx.Response(200, json={"data": []})
        ),
    )
    first_app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-first.db'}"),
        llm_client=first_primary,
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )
    with TestClient(first_app) as client:
        updated = client.put(
            "/api/settings",
            json={
                "resume_llm": {
                    "base_url": "http://resume.test/v1",
                    "model": "resume",
                }
            },
        )
        assert updated.status_code == 200

    requests = []

    def restarted_handler(request):
        requests.append(str(request.url))
        return httpx.Response(200, json={"data": [{"id": "persisted-resume"}]})

    second_primary = LocalLLMClient(
        base_url="http://main.test/v1",
        transport=httpx.MockTransport(restarted_handler),
    )
    second_app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-second.db'}"),
        llm_client=second_primary,
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )
    with TestClient(second_app) as client:
        dedicated = client.app.state.resume_llm_client
        assert dedicated is not second_primary
        result = client.portal.call(dedicated.probe)

        assert result["models"] == ["persisted-resume"]
        assert requests == ["http://resume.test/v1/models"]


def test_candidate_prep_api_returns_structured_workflow(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'workflow-api.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        response = client.post(
            "/api/workflows/candidate-prep",
            json={
                "session_id": session_id,
                "resume_text": "Python engineer",
                "job_description": "Platform engineer",
                "company_name": "Example",
            },
        )
    assert response.status_code == 200
    state = response.json()["state"]
    assert state["workflow"]["status"] == "completed"
    assert state["strategy"]["summary"] == "Show impact"
    first_question = state["mock_interview"]["questions"][0]
    assert first_question["question"].startswith("Design it")
    assert "结果如何验证" in first_question["question"]
    assert first_question["question_requirements"]
    assert first_question["example_answer"]


def test_job_requirement_review_confirms_edits_and_deletes_inferred_items(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'job-review.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        analyzed = client.post(
            "/api/analysis/job",
            json={"session_id": session_id, "text": "高级平台工程师"},
        )
        confirmed = client.patch(
            f"/api/analysis/job/{session_id}/requirements/0",
            json={"action": "confirm"},
        )
        edited = client.patch(
            f"/api/analysis/job/{session_id}/requirements/0",
            json={"action": "edit", "text": "需要设计高可用平台架构"},
        )
        deleted = client.patch(
            f"/api/analysis/job/{session_id}/requirements/0",
            json={"action": "delete"},
        )

    assert analyzed.status_code == 200
    assert analyzed.json()["state"]["job_review"]["requirements"][0]["origin"] == "inferred"
    assert confirmed.status_code == 200
    assert confirmed.json()["state"]["job_review"]["requirements"][0]["origin"] == "explicit"
    assert edited.status_code == 200
    requirement = edited.json()["state"]["job_review"]["requirements"][0]
    assert requirement == {"text": "需要设计高可用平台架构", "origin": "explicit"}
    assert deleted.status_code == 200
    assert deleted.json()["state"]["job_review"]["requirements"] == []


def test_mock_interview_api_progresses_to_completion(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'mock-api.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/workflows/candidate-prep",
            json={
                "session_id": session_id,
                "resume_text": "Python engineer",
                "job_description": "Platform engineer",
                "company_name": "Example",
            },
        )
        started = client.post(f"/api/mock-interviews/{session_id}/start")
        question_id = started.json()["current_question"]["id"]
        assert started.json()["current_question"]["answer_framework"] is not None
        answered = client.post(
            f"/api/mock-interviews/{session_id}/answers",
            json={"question_id": question_id, "answer": "I explained the design trade-offs."},
        )
    assert answered.status_code == 200
    # The interview does not auto-complete after one answer.
    assert answered.json()["mock_session"]["status"] == "active"
    assert answered.json()["current_question"] is not None
    assert answered.json()["answered_questions"] == 1


def test_evaluation_api_generates_dual_side_report(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'evaluation-api.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        without_evidence = client.post(f"/api/evaluations/{session_id}")
        assert without_evidence.status_code == 409

        client.post(
            "/api/workflows/candidate-prep",
            json={
                "session_id": session_id,
                "resume_text": "Python",
                "job_description": "Platform",
                "company_name": "Example",
            },
        )
        started = client.post(f"/api/mock-interviews/{session_id}/start").json()
        client.post(
            f"/api/mock-interviews/{session_id}/answers",
            json={"question_id": started["current_question"]["id"], "answer": "Clear design"},
        )
        client.post(f"/api/mock-interviews/{session_id}/finish")
        response = client.post(f"/api/evaluations/{session_id}")
    assert response.status_code == 200
    assert response.json()["state"]["evaluation"]["recommendation"] == "insufficient_evidence"
    assert "证据不足" in response.json()["state"]["feedback"]["overall"]
    notes = response.json()["state"]["feedback"]["interviewer_notes"]
    assert notes[:2] == [
        "当前仅有 1 条证据，覆盖 1 个胜任力；未达到招聘决策门槛。",
        "证据不足不是负面证据，需要继续采集独立回答",
    ]


def test_human_answer_review_overrides_score_and_classifies_linked_evidence(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'human-review-api.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/workflows/candidate-prep",
            json={
                "session_id": session_id,
                "resume_text": "Python",
                "job_description": "Platform",
                "company_name": "Example",
            },
        )
        started = client.post(f"/api/mock-interviews/{session_id}/start").json()
        answered = client.post(
            f"/api/mock-interviews/{session_id}/answers",
            json={"question_id": started["current_question"]["id"], "answer": "Clear design"},
        ).json()
        record_id = answered["mock_session"]["responses"][0]["id"]

        response = client.patch(
            f"/api/evaluations/{session_id}/answers/{record_id}/review",
            json={
                "content": 0.9,
                "technical_depth": 0.8,
                "structure": 0.7,
                "impact": 0.6,
                "evidence_polarity": "negative",
                "note": "已核对原始回答",
            },
        )

    assert response.status_code == 200
    state = response.json()["state"]
    evaluation = state["mock_session"]["responses"][0]["evaluation"]
    assert evaluation["scoring_source"] == "human"
    assert evaluation["review_status"] == "reviewed"
    assert evaluation["content"] == pytest.approx(0.9)
    assert state["evidence"][0]["polarity"] == "negative"
    assert state["evidence"][0]["confidence"] == pytest.approx(0.75)
    assert state["evaluation"]["finalized_at"] is None


def test_mock_interview_asks_followup_before_advancing(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'followup.db'}")
    app = create_app(storage=storage, llm_client=FollowupWorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/workflows/candidate-prep",
            json={
                "session_id": session_id,
                "resume_text": "Python",
                "job_description": "Platform",
                "company_name": "Example",
            },
        )
        started = client.post(f"/api/mock-interviews/{session_id}/start").json()
        question_id = started["current_question"]["id"]
        first = client.post(
            f"/api/mock-interviews/{session_id}/answers",
            json={"question_id": question_id, "answer": "I made a trade-off."},
        ).json()
        assert first["mock_session"]["status"] == "active"
        # Advancing chooses a branch from the answer's auditable coverage gaps.
        advanced = client.post(f"/api/mock-interviews/{session_id}/next").json()
        assert advanced["mock_session"]["pending_follow_up"]
        assert advanced["mock_session"]["pending_follow_up_stage"] in {
            "foundation",
            "evidence",
            "tradeoff",
            "pressure",
        }
        assert advanced["mock_session"]["pending_follow_up_rationale"]
        assert (
            advanced["current_question"]["question"]
            == advanced["mock_session"]["pending_follow_up"]
        )
        follow_id = advanced["current_question"]["id"]
        second = client.post(
            f"/api/mock-interviews/{session_id}/answers",
            json={"question_id": follow_id, "answer": "Latency fell by 30%."},
        ).json()
        finished = client.post(f"/api/mock-interviews/{session_id}/finish").json()
    assert second["mock_session"]["status"] == "active"
    assert finished["mock_session"]["status"] == "completed"
    assert finished["mock_session"]["pending_follow_up"] == ""
    assert finished["mock_session"]["pending_follow_up_stage"] == ""
    assert second["mock_session"]["responses"][1]["is_follow_up"] is True
    assert second["mock_session"]["responses"][1]["follow_up_stage"]


def test_interviewer_transcript_import_generates_hiring_report(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'transcript.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        response = client.post(
            f"/api/evaluations/{session_id}/transcript",
            json={
                "entries": [
                    {
                        "competency": "System Design",
                        "question": "Design a service",
                        "answer": "I measured latency and explained the trade-offs.",
                    }
                ]
            },
        )
    assert response.status_code == 200
    result = response.json()["state"]
    assert len(result["live_interview_records"]) == 1
    assert result["evidence"][0]["source"] == "live_interview"
    assert result["evaluation"]["recommendation"] == "insufficient_evidence"
    assert "不能给出录用" in result["feedback"]["recommendation_reasoning"]


def test_live_audio_transcription_and_next_question_flow(tmp_path):
    def asr_handler(request):
        assert request.url.path == "/v1/audio/transcriptions"
        return httpx.Response(
            200,
            json={"text": "我使用压测比较两个方案，最终 P95 降低了 40%。"},
        )

    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-interview.db'}")
    asr = ASRClient(
        base_url="http://asr.test:9001",
        transport=httpx.MockTransport(asr_handler),
    )
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        asr_client=asr,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/analysis/job",
            json={"session_id": session_id, "text": "System Design Engineer"},
        )
        denied = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": False, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        started = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        recording_id = _register_live_capture(client, session_id)
        transcribed = client.post(
            f"/api/live-interviews/{session_id}/audio",
            data={
                "speaker": "candidate",
                "language": "zh",
                "recording_id": recording_id,
            },
            files={"file": ("answer.webm", b"fake-audio", "audio/webm")},
        )
        planned = client.post(f"/api/live-interviews/{session_id}/suggestions")
        suggestion = planned.json()["state"]["live_interview"]["suggestions"][0]
        adopted = client.patch(
            f"/api/live-interviews/{session_id}/suggestions/{suggestion['id']}",
            json={"status": "adopted"},
        )
        answered = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={
                "speaker": "candidate",
                "text": "我用压测和灰度验证方案，P95 延迟下降 30%，异常时可以按灰度批次回滚。",
            },
        )
        answer_segment = answered.json()["state"]["live_interview"]["segments"][-1]
        question_segment = adopted.json()["state"]["live_interview"]["segments"][-1]
        corrected = client.patch(
            f"/api/live-interviews/{session_id}/segments/{answer_segment['id']}",
            json={
                "speaker": "candidate",
                "text": "我用压测和灰度验证方案，P95 延迟下降 40%，异常时可以按灰度批次回滚。",
            },
        )
        confirmed = client.post(
            f"/api/live-interviews/{session_id}/segments/{answer_segment['id']}/evidence",
            json={
                "question_segment_id": question_segment["id"],
                "competency": "System Design",
            },
        )
        duplicate = client.post(
            f"/api/live-interviews/{session_id}/segments/{answer_segment['id']}/evidence",
            json={"competency": "System Design"},
        )
        edit_confirmed = client.patch(
            f"/api/live-interviews/{session_id}/segments/{answer_segment['id']}",
            json={"text": "确认后不应该直接改写"},
        )

    assert denied.status_code == 409
    assert started.status_code == 200
    segment = transcribed.json()["state"]["live_interview"]["segments"][0]
    assert segment["source"] == "asr"
    assert "P95" in segment["text"]
    assert planned.status_code == 200
    assert suggestion["competency"] == "System Design"
    assert suggestion["question_type"] == "follow_up"
    final_live = adopted.json()["state"]["live_interview"]
    assert final_live["suggestions"][0]["status"] == "adopted"
    assert final_live["segments"][-1]["speaker"] == "interviewer"
    assert corrected.status_code == 200
    corrected_segment = corrected.json()["state"]["live_interview"]["segments"][-1]
    assert "40%" in corrected_segment["text"]
    assert confirmed.status_code == 200
    confirmed_state = confirmed.json()["state"]
    assert confirmed_state["live_interview_records"][0]["source"] == "live_interview"
    assert "40%" in confirmed_state["live_interview_records"][0]["answer"]
    assert confirmed_state["live_interview_records"][0]["transcript_segment_ids"] == [
        question_segment["id"],
        answer_segment["id"],
    ]
    assert confirmed_state["evidence"][0]["source"] == "live_interview"
    assert confirmed_state["evidence"][0]["competency"] == "System Design"
    assert duplicate.status_code == 409
    assert edit_confirmed.status_code == 409


def test_live_review_queue_batch_confirms_pending_candidate_answers(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-batch.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/analysis/job",
            json={"session_id": session_id, "text": "System Design Engineer"},
        )
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "interviewer", "text": "请讲一次架构权衡。"},
        )
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我比较了缓存和数据库扩容。"},
        )
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我用灰度发布验证并降低了延迟。"},
        )
        confirmed = client.post(
            f"/api/live-interviews/{session_id}/evidence/batch",
            json={"competency": "System Design"},
        )
        empty = client.post(
            f"/api/live-interviews/{session_id}/evidence/batch",
            json={"competency": "System Design"},
        )

    assert confirmed.status_code == 200
    state = confirmed.json()["state"]
    assert len(state["live_interview_records"]) == 2
    assert len(state["evidence"]) == 2
    assert {item["source"] for item in state["evidence"]} == {"live_interview"}
    assert empty.status_code == 409


def test_live_evidence_can_be_revoked_and_reconfirmed(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-revoke.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        question = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "interviewer", "text": "请讲一次线上故障处理。"},
        ).json()["state"]["live_interview"]["segments"][0]
        answer = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我先止血，再复盘根因并补监控。"},
        ).json()["state"]["live_interview"]["segments"][1]
        confirmed = client.post(
            f"/api/live-interviews/{session_id}/segments/{answer['id']}/evidence",
            json={"question_segment_id": question["id"], "competency": "Incident Response"},
        )
        confirmed_state = confirmed.json()["state"]
        record_id = confirmed_state["live_interview_records"][0]["id"]

        revoked = client.delete(f"/api/live-interviews/{session_id}/evidence/{record_id}")
        reconfirmed = client.post(
            f"/api/live-interviews/{session_id}/segments/{answer['id']}/evidence",
            json={"question_segment_id": question["id"], "competency": "Incident Response"},
        )

    assert confirmed.status_code == 200
    assert confirmed_state["evidence"][0]["source_record_id"] == record_id
    assert revoked.status_code == 200
    revoked_state = revoked.json()["state"]
    assert revoked_state["live_interview_records"] == []
    assert revoked_state["evidence"] == []
    assert reconfirmed.status_code == 200
    reconfirmed_state = reconfirmed.json()["state"]
    assert len(reconfirmed_state["live_interview_records"]) == 1
    assert len(reconfirmed_state["evidence"]) == 1
    assert reconfirmed_state["live_interview_records"][0]["id"] != record_id


def test_live_candidate_segments_can_be_merged_into_one_evidence_record(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-merge.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        question = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "interviewer", "text": "请讲一次架构演进。"},
        ).json()["state"]["live_interview"]["segments"][0]
        first_text = "第一阶段我先拆分核心服务，把订单、库存和支付链路分开，并补了基础压测。"
        second_text = "第二阶段我补了监控、灰度发布和回滚预案，最终上线后错误率和延迟指标都比较稳定，以上。"
        first = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": first_text},
        ).json()["state"]["live_interview"]["segments"][1]
        second = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": second_text},
        ).json()["state"]["live_interview"]["segments"][2]
        live = client.get(f"/api/live-interviews/{session_id}").json()["state"]["live_interview"]
        boundary = live["answer_boundary_suggestions"][0]
        assert boundary["question_segment_id"] == question["id"]
        assert boundary["answer_segment_ids"] == [first["id"], second["id"]]
        assert boundary["confidence"] >= 0.75
        assert "已关联最近面试官问题" in boundary["confidence_factors"]
        assert "末段出现回答结束信号" in boundary["confidence_factors"]
        assert live["action_card"]["action_type"] == "merge_boundary"
        assert live["action_card"]["primary_cta"] == "按边界建议合并"

        merged = client.post(
            f"/api/live-interviews/{session_id}/evidence/merge",
            json={
                "segment_ids": [second["id"], first["id"]],
                "question_segment_id": question["id"],
                "competency": "Architecture Evolution",
            },
        )
        duplicate = client.post(
            f"/api/live-interviews/{session_id}/segments/{first['id']}/evidence",
            json={"question_segment_id": question["id"], "competency": "Architecture Evolution"},
        )

    assert merged.status_code == 200
    state = merged.json()["state"]
    assert len(state["live_interview_records"]) == 1
    assert len(state["evidence"]) == 1
    record = state["live_interview_records"][0]
    assert record["answer"] == f"{first_text}\n{second_text}"
    assert record["transcript_segment_ids"] == [question["id"], first["id"], second["id"]]
    assert state["live_interview"]["answer_boundary_suggestions"] == []
    # After merging evidence, the next-question planner auto-runs: a pending
    # suggestion takes priority, so the card moves straight to deciding it.
    assert state["live_interview"]["action_card"]["action_type"] == "decide_question"
    assert state["evidence"][0]["source_record_id"] == record["id"]
    assert duplicate.status_code == 409


def test_live_transcript_dedupes_and_rolls_context_for_question_planning(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-context.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "interviewer", "text": "请讲一次系统设计经历。"},
        )
        first = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我先梳理瓶颈并设计分层缓存。"},
        ).json()["state"]["live_interview"]
        duplicate = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我先梳理瓶颈并设计分层缓存。"},
        ).json()["state"]["live_interview"]
        for index in range(12):
            client.post(
                f"/api/live-interviews/{session_id}/segments",
                json={
                    "speaker": "candidate",
                    "text": f"补充第 {index} 点：我用指标验证方案并记录风险。",
                },
            )
        planned = client.post(f"/api/live-interviews/{session_id}/suggestions")
        events = client.get("/api/debug/events").json()["events"]

    assert len(first["segments"]) == 2
    assert len(duplicate["segments"]) == 2
    assert duplicate["duplicate_segments_dropped"] == 1
    live = planned.json()["state"]["live_interview"]
    assert live["summarized_until_sequence"] > 0
    assert "请讲一次系统设计经历" in live["rolling_summary"]
    planned_events = [item for item in events if item["action"] == "live_question_planned"]
    assert planned_events
    assert "summary_until=" in planned_events[0]["detail"]
    assert "duplicates_dropped=1" in planned_events[0]["detail"]
    assert "action=merge_boundary" in planned_events[0]["detail"]


def test_live_coverage_guidance_tracks_evidence_gaps(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-guidance.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/analysis/job",
            json={"session_id": session_id, "text": "System Design Engineer"},
        )
        started = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        ).json()["state"]["live_interview"]
        question = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "interviewer", "text": "请讲一次架构权衡。"},
        ).json()["state"]["live_interview"]["segments"][0]
        first = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我通过压测比较方案并灰度上线。"},
        ).json()["state"]["live_interview"]["segments"][1]
        one_evidence = client.post(
            f"/api/live-interviews/{session_id}/segments/{first['id']}/evidence",
            json={"question_segment_id": question["id"], "competency": "System Design"},
        ).json()["state"]["live_interview"]["coverage_guidance"][0]
        second = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我补充了容量评估、回滚预案和故障复盘。"},
        ).json()["state"]["live_interview"]["segments"][-1]
        two_evidence = client.post(
            f"/api/live-interviews/{session_id}/segments/{second['id']}/evidence",
            json={"question_segment_id": question["id"], "competency": "System Design"},
        ).json()["state"]["live_interview"]["coverage_guidance"][0]
        client.post(f"/api/live-interviews/{session_id}/suggestions")
        events = client.get("/api/debug/events").json()["events"]

    initial = started["coverage_guidance"][0]
    assert initial["competency"] == "System Design"
    assert initial["priority"] == "high"
    assert initial["evidence_count"] == 0
    assert one_evidence["priority"] == "medium"
    assert one_evidence["evidence_count"] == 1
    # Two low-evidence answers still need follow-up: count alone cannot erase the
    # missing case/result signals exposed by score calibration.
    assert two_evidence["priority"] == "medium"
    assert two_evidence["evidence_count"] == 2
    planned_events = [item for item in events if item["action"] == "live_question_planned"]
    assert planned_events
    assert "top_gap=System Design" in planned_events[0]["detail"]
    assert "top_boundary_confidence=" in planned_events[0]["detail"]


def test_live_question_usage_tracks_blueprint_progress(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-question-usage.db'}")
    app = create_app(
        storage=storage,
        llm_client=BlueprintFallbackQuestionLLM(),
        configure_llm=False,
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        designed = client.post(
            "/api/workflows/enterprise-design",
            json={
                "session_id": session_id,
                "resume_text": "Platform engineer with reliability experience.",
                "job_description": "平台工程师\n岗位职责：系统设计、稳定性治理。\n任职要求：事故复盘。",
                "company_name": "Example",
            },
        ).json()["state"]
        started = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        ).json()["state"]["live_interview"]
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我做过一次缓存架构演进。"},
        )
        planned = client.post(f"/api/live-interviews/{session_id}/suggestions").json()["state"]
        suggestion = planned["live_interview"]["suggestions"][0]
        adopted = client.patch(
            f"/api/live-interviews/{session_id}/suggestions/{suggestion['id']}",
            json={"status": "adopted"},
        ).json()["state"]["live_interview"]
        events = client.get("/api/debug/events").json()["events"]

    assert len(designed["blueprint"]["rounds"][0]["questions"]) == 2
    assert {item["status"] for item in started["question_usage"]} == {"pending"}
    assert [item["status"] for item in planned["live_interview"]["question_usage"]].count(
        "suggested"
    ) == 1
    assert suggestion["source_question_id"]
    assert adopted["used_question_ids"] == [suggestion["source_question_id"]]
    assert adopted["question_usage"][0]["status"] == "pending"
    assert adopted["question_usage"][-1]["status"] == "used"
    assert adopted["question_usage"][-1]["suggested_count"] == 1
    assert planned["live_interview"]["action_card"]["action_type"] == "confirm_evidence"
    assert planned["live_interview"]["action_card"]["primary_cta"] == "确认为证据"
    planned_events = [item for item in events if item["action"] == "live_question_planned"]
    assert planned_events
    assert "pending_blueprint=1" in planned_events[0]["detail"]
    assert "action=confirm_evidence" in planned_events[0]["detail"]


def test_live_evidence_can_be_reevaluated_with_updated_competency(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-reevaluate.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        question = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "interviewer", "text": "请讲一次跨团队项目。"},
        ).json()["state"]["live_interview"]["segments"][0]
        answer = client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我协调平台和业务团队完成迁移。"},
        ).json()["state"]["live_interview"]["segments"][1]
        confirmed = client.post(
            f"/api/live-interviews/{session_id}/segments/{answer['id']}/evidence",
            json={"question_segment_id": question["id"], "competency": "System Design"},
        ).json()["state"]
        record_id = confirmed["live_interview_records"][0]["id"]

        reevaluated = client.patch(
            f"/api/live-interviews/{session_id}/evidence/{record_id}/reevaluate",
            json={"competency": "Cross-team Collaboration"},
        )

    assert reevaluated.status_code == 200
    state = reevaluated.json()["state"]
    assert len(state["live_interview_records"]) == 1
    assert len(state["evidence"]) == 1
    record = state["live_interview_records"][0]
    evidence = state["evidence"][0]
    assert record["id"] == record_id
    assert record["competency"] == "Cross-team Collaboration"
    assert evidence["competency"] == "Cross-team Collaboration"
    assert evidence["source_record_id"] == record_id
    assert record["transcript_segment_ids"] == [question["id"], answer["id"]]


def test_fact_card_decision_api_updates_session_state(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'facts.db'}")
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        search_provider=FactSearchProvider(),
        configure_llm=False,
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        analyzed = client.post(
            "/api/analysis/company",
            json={"session_id": session_id, "name": "Example", "context": "Example develops AI systems."},
        ).json()["state"]
        card_id = analyzed["fact_cards"][0]["id"]

        accepted = client.patch(
            f"/api/intelligence/{session_id}/facts/{card_id}",
            json={"action": "accept", "note": "用户确认"},
        )
        rejected = client.patch(
            f"/api/intelligence/{session_id}/facts/{card_id}",
            json={"action": "reject", "note": "来源冲突"},
        )
        reset = client.patch(
            f"/api/intelligence/{session_id}/facts/{card_id}",
            json={"action": "reset"},
        )

    assert accepted.status_code == 200
    assert accepted.json()["state"]["fact_cards"][0]["status"] == "accepted"
    assert accepted.json()["state"]["fact_cards"][0]["note"] == "用户确认"
    assert rejected.status_code == 200
    assert rejected.json()["state"]["fact_cards"][0]["status"] == "rejected"
    assert reset.status_code == 200
    assert reset.json()["state"]["fact_cards"][0]["status"] == "inferred"


def test_live_interviewer_workflow_reaches_sufficient_evidence_evaluation(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-sufficient.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/workflows/enterprise-design",
            json={
                "session_id": session_id,
                "resume_text": "Platform engineer with stability and architecture experience.",
                "job_description": "平台工程师\n岗位职责：架构设计、稳定性治理。\n任职要求：系统设计、故障复盘。\n团队背景：平台团队。",
                "company_name": "Example",
            },
        )
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        entries = [
            ("interviewer", "请讲一次架构权衡。"),
            ("candidate", "我比较了缓存治理、数据库扩容和限流，最终选择先治理缓存。"),
            ("candidate", "上线前通过压测验证 P95 延迟，并保留灰度回滚方案。"),
            ("candidate", "事故后我组织复盘，补充告警、演练和 runbook。"),
        ]
        for speaker, text in entries:
            client.post(
                f"/api/live-interviews/{session_id}/segments",
                json={"speaker": speaker, "text": text},
            )
        state = client.get(f"/api/live-interviews/{session_id}").json()["state"]
        candidates = [
            item for item in state["live_interview"]["segments"] if item["speaker"] == "candidate"
        ]
        client.post(
            f"/api/live-interviews/{session_id}/segments/{candidates[0]['id']}/evidence",
            json={"competency": "System Design"},
        )
        client.post(
            f"/api/live-interviews/{session_id}/segments/{candidates[1]['id']}/evidence",
            json={"competency": "System Design"},
        )
        client.post(
            f"/api/live-interviews/{session_id}/segments/{candidates[2]['id']}/evidence",
            json={"competency": "Incident Review"},
        )
        client.post(
            f"/api/live-interviews/{session_id}/status",
            json={"status": "completed", "expected_revision": 1},
        )
        evaluated = client.post(f"/api/evaluations/{session_id}")

    assert evaluated.status_code == 200
    result = evaluated.json()["state"]
    assert len(result["evidence"]) == 3
    assert {item["competency"] for item in result["evidence"]} == {
        "System Design",
        "Incident Review",
    }
    assert result["evaluation"]["recommendation"] != "insufficient_evidence"


def test_live_status_revision_barrier_blocks_late_mutations(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'live-status-cas.db'}")
    app = create_app(storage=storage, llm_client=None, configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]

        barrier = client.post(
            f"/api/live-interviews/{session_id}/status-barrier",
            json={"expected_revision": 0},
        )
        stale_start = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        current = client.get(f"/api/live-interviews/{session_id}")

        assert barrier.status_code == 200
        assert barrier.json()["state"]["live_interview"]["status_revision"] == 1
        assert stale_start.status_code == 409
        assert current.json()["state"]["live_interview"]["status"] == "idle"

        started = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 1, "operation_id": str(uuid4())},
        )
        assert started.status_code == 200
        assert started.json()["state"]["live_interview"]["status_revision"] == 2


def test_long_live_interview_rolls_and_dedupes(tmp_path):
    """Long interview keeps context bounded: summary activates, duplicates drop,
    and the next-question planner stays fast enough for the 5s product target."""
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'long-live.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        # A real long interview: 30+ alternating Q&A turns, with a few
        # consecutive candidate chunks (as chunked ASR would produce).
        for index in range(30):
            client.post(
                f"/api/live-interviews/{session_id}/segments",
                json={"speaker": "interviewer", "text": f"第 {index} 轮：请补充一次架构决策。"},
            )
            client.post(
                f"/api/live-interviews/{session_id}/segments",
                json={"speaker": "candidate", "text": f"第 {index} 轮回答：我对比方案并用压测验证。"},
            )
            if index % 10 == 9:
                # Chunked ASR split: a second candidate chunk for the same answer.
                client.post(
                    f"/api/live-interviews/{session_id}/segments",
                    json={"speaker": "candidate", "text": f"第 {index} 轮补充：P95 延迟下降，以上就是。"},
                )
        # Duplicate segment is dropped without growing the transcript. The
        # dedup window only covers the most recent LIVE_RECENT_SEGMENT_WINDOW
        # segments, so re-send the last candidate turn.
        before = client.get(f"/api/live-interviews/{session_id}").json()["state"]["live_interview"]
        duplicate_count = before["duplicate_segments_dropped"]
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "第 29 轮补充：P95 延迟下降，以上就是。"},
        )
        after = client.get(f"/api/live-interviews/{session_id}").json()["state"]["live_interview"]
        assert after["duplicate_segments_dropped"] == duplicate_count + 1

        # Rolling summary must be active and the raw transcript must not have
        # grown proportionally with the summary.
        planned = client.post(f"/api/live-interviews/{session_id}/suggestions")
        planned_json = planned.json()["state"]["live_interview"]
        assert planned_json["summarized_until_sequence"] > 0
        assert planned_json["rolling_summary"]
        assert planned_json["duplicate_segments_dropped"] >= 1
        # Next-question planning must stay within the 5s product budget.
        assert planned.elapsed.total_seconds() <= 5.0
        # Boundary suggestions exist for merged candidate runs.
        assert planned_json["answer_boundary_suggestions"]



class _FakeOmniClient:
    def __init__(self):
        self.mode = "asr_text"
        self.calls = []

    async def suggest_next_question(self, audio_bytes, *, content_type="audio/wav", context="", stream=True):
        self.calls.append({"bytes": len(audio_bytes), "content_type": content_type, "context": context})
        return "请再补充说明一下方案权衡与量化结果。"

    async def probe_capability(self):
        return {"ok": True, "latency_ms": 1200.0, "sample": "测试建议"}

    def configure(self, **kwargs):
        if kwargs.get("mode"):
            self.mode = kwargs["mode"]

    def status(self):
        return {"enabled": True, "mode": self.mode, "name": "测试直连", "base_url": "http://omni.test", "model": "m", "api_key_configured": False, "capability": {}}

    def secret_snapshot(self):
        return {"mode": self.mode, "name": "测试直连", "base_url": "http://omni.test", "api_key": "", "model": "m"}

    async def close(self):
        pass


def test_live_audio_direct_mode_injects_suggestion(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'audio-direct.db'}")
    omni = _FakeOmniClient()
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        omni_client=omni,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        # Switch to audio-direct via settings
        switched = client.put(
            "/api/settings",
            json={"live_audio": {"mode": "audio_direct", "name": "我的直连"}, "persist": False},
        )
        assert switched.status_code == 200
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        # Add an interviewer question + candidate answer so the audio-direct
        # context has transcript material to include.
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "interviewer", "text": "请讲一次系统设计经历。"},
        )
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我对比了缓存和数据库方案。"},
        )
        recording_id = _register_live_capture(client, session_id)
        uploaded = client.post(
            f"/api/live-interviews/{session_id}/audio",
            data={
                "speaker": "candidate",
                "language": "zh",
                "recording_id": recording_id,
            },
            files={"file": ("answer.wav", b"fake-audio-bytes", "audio/wav")},
        )
        assert uploaded.status_code == 200
        live = uploaded.json()["state"]["live_interview"]
        # Audio-direct creates no new transcript segment (the model hears the
        # audio directly); only the manually-added candidate segment remains.
        candidate_segments = [s for s in live["segments"] if s["speaker"] == "candidate"]
        assert len(candidate_segments) == 1  # only the manual one
        assert not any(s["source"] == "asr" for s in live["segments"])
        assert live["suggestions"], "audio-direct should inject a suggestion"
        assert live["suggestions"][-1]["suggested_question"] == "请再补充说明一下方案权衡与量化结果。"
        # The omni client must have received bounded context with transcript.
        assert omni.calls, "omni client should have been called"
        sent_context = omni.calls[-1]["context"]
        assert "岗位能力" in sent_context
        assert "请讲一次系统设计经历" in sent_context
        assert len(sent_context) <= 2600  # bounded by OMNI_CONTEXT_CHAR_LIMIT


def test_settings_probe_live_audio_reports_capability(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'probe.db'}")
    omni = _FakeOmniClient()
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        omni_client=omni,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        probe = client.post("/api/settings/probe-live-audio")
        assert probe.status_code == 200
        assert probe.json()["capability"]["ok"] is True


class _StructuringLLM:
    async def chat(self, messages, **kwargs):
        return '{"education":[{"category":"education","institution":"SMIC","title":"High School","date_range":"08/2020 - 06/2024"}],"employment":[{"category":"employment","institution":"ZUORA","title":"Senior Recruiting Manager","date_range":"2018-07 - 2022-12"}]}'

    async def embed(self, text):
        return []


def test_resume_upload_with_llm_structure(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-llm.db'}")
    llm = _StructuringLLM()
    app = create_app(
        storage=storage,
        llm_client=llm,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
        resume_llm_client=llm,
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        uploaded = client.post(
            f"/api/resumes/{session_id}/upload",
            data={"structure": "llm"},
            files={"file": ("resume.pdf", make_resume_pdf(), "application/pdf")},
        )
        assert uploaded.status_code == 200
        review = uploaded.json()["state"]["resume_review"]
        assert review["structured_by"] == "llm"
        assert review["structured"], "LLM structuring should populate sections"
        categories = {s["category"] for s in review["structured"]}
        assert categories == {"education", "employment"}


def test_resume_upload_default_rules_structure(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-rules.db'}")
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        uploaded = client.post(
            f"/api/resumes/{session_id}/upload",
            files={"file": ("resume.pdf", make_resume_pdf(), "application/pdf")},
        )
        assert uploaded.status_code == 200
        review = uploaded.json()["state"]["resume_review"]
        assert review["structured_by"] == "rules"
        assert review["structured"] == []


def test_scanned_resume_upload_uses_local_ocr_and_marks_text_unverified(
    tmp_path, monkeypatch
):
    from interview_os.services import resume_ocr

    recognized = (
        "周岚\n供应链运营经理\n工作经历\n海岳消费品 | 供应链运营经理 | 2020.06 - 至今\n"
        "将缺货率从 8.1% 降至 3.4%\n技能\n供应链规划\n教育经历\n南方管理学院"
    )
    monkeypatch.setattr(resume_ocr, "ocr_pages_with_vision", lambda images: recognized)
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'resume-ocr.db'}")
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )

    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        uploaded = client.post(
            f"/api/resumes/{session_id}/upload",
            files={"file": ("scan.pdf", make_scanned_resume_pdf(), "application/pdf")},
        )

    assert uploaded.status_code == 200
    state = uploaded.json()["state"]
    assert "周岚" in state["candidate"]["raw_resume_text"]
    issues = state["resume_review"]["issues"]
    assert issues[0]["code"] == "ocr_transcription_unverified"


def test_suggestions_stream_requires_active_live(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'stream-guard.db'}")
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        # Live not started -> streaming must refuse instead of calling the LLM.
        resp = client.post(f"/api/live-interviews/{session_id}/suggestions/stream")
        assert resp.status_code == 409


def test_suggestions_stream_json_encodes_sse_data(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'stream-json.db'}")
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        started = client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        assert started.status_code == 200
        client.post(
            f"/api/live-interviews/{session_id}/segments",
            json={"speaker": "candidate", "text": "我说明了方案取舍。"},
        )
        response = client.post(f"/api/live-interviews/{session_id}/suggestions/stream")
        assert response.status_code == 200
        data_lines = [line for line in response.text.splitlines() if line.startswith("data: ")]
        assert data_lines
        events = [json.loads(line[6:]) for line in data_lines]
        assert all(event["type"] in {"append", "replace"} for event in events)
        assert all(isinstance(event["text"], str) for event in events)


class _PartialFailingStreamLLM(WorkflowLLM):
    async def chat_stream(self, messages, **kwargs):
        yield "半截内部输出"
        raise RuntimeError("secret backend detail")


def test_suggestions_stream_replaces_partial_output_after_failure(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'stream-fallback.db'}")
    app = create_app(
        storage=storage,
        llm_client=_PartialFailingStreamLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            f"/api/live-interviews/{session_id}/start",
            json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())},
        )
        response = client.post(f"/api/live-interviews/{session_id}/suggestions/stream")
        events = [
            json.loads(line[6:])
            for line in response.text.splitlines()
            if line.startswith("data: ")
        ]
        assert events[0] == {"type": "append", "text": "半截内部输出"}
        assert events[-1]["type"] == "replace"
        assert "secret backend detail" not in response.text
        state = client.get(f"/api/live-interviews/{session_id}").json()["state"]
        persisted = state["live_interview"]["suggestions"][-1]["suggested_question"]
        assert persisted == events[-1]["text"]
        assert "半截内部输出" not in persisted


class _DiarizeOmni:
    """Stub omni client returning a two-speaker dialog split."""

    async def transcribe_diarize(self, audio_bytes, *, content_type="audio/wav"):
        return [
            {"speaker": "interviewer", "text": "请讲一次系统设计。"},
            {"speaker": "candidate", "text": "我对比了缓存和数据库方案。"},
            {"speaker": "candidate", "text": "用压测验证性能提升。"},
        ]

    def configure(self, **kwargs):
        pass

    async def suggest_next_question(self, audio_bytes, *, content_type="audio/wav", context="", stream=True):
        return "追问建议"

    async def probe_capability(self):
        return {"ok": True}

    def status(self):
        return {"enabled": True}

    def secret_snapshot(self):
        return {}

    async def close(self):
        pass


def test_audio_direct_dialogue_mode_splits_speakers(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'dialogue.db'}")
    omni = _DiarizeOmni()
    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        omni_client=omni,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        client.put("/api/settings", json={"live_audio": {"mode": "audio_direct"}, "persist": False})
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(f"/api/live-interviews/{session_id}/start", json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())})
        recording_id = _register_live_capture(client, session_id)
        uploaded = client.post(
            f"/api/live-interviews/{session_id}/audio",
            data={
                "mode": "dialogue",
                "speaker": "unknown",
                "language": "zh",
                "recording_id": recording_id,
            },
            files={"file": ("dialog.wav", b"audio-bytes", "audio/wav")},
        )
        assert uploaded.status_code == 200
        live = uploaded.json()["state"]["live_interview"]
        segments = live["segments"]
        speakers = [s["speaker"] for s in segments]
        # Interviewer + one merged candidate segment (two candidate utterances merged).
        assert speakers[0] == "interviewer"
        assert speakers[1] == "candidate"
        assert len(segments) == 2
        assert "用压测验证性能提升" in segments[1]["text"]


def test_audio_direct_dialogue_mode_single_utterance_creates_segment(tmp_path):
    """A single finalized utterance in audio_direct dialogue mode becomes one segment.

    This is the contract the silence-based continuous listening path relies on:
    the browser finalizes a whole utterance at ~0.6s of silence and uploads it with
    mode=dialogue, and the omni diarize produces exactly one transcript segment.
    """
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'utterance.db'}")

    class _SingleUtteranceOmni:
        async def transcribe_diarize(self, audio_bytes, *, content_type="audio/wav"):
            return [{"speaker": "candidate", "text": "我对比了缓存和数据库方案，用压测验证性能提升。"}]

        def configure(self, **kwargs):
            pass

        async def suggest_next_question(self, audio_bytes, *, content_type="audio/wav", context="", stream=True):
            return "追问建议"

        async def probe_capability(self):
            return {"ok": True}

        def status(self):
            return {"enabled": True}

        def secret_snapshot(self):
            return {}

        async def close(self):
            pass

    app = create_app(
        storage=storage,
        llm_client=WorkflowLLM(),
        configure_llm=False,
        omni_client=_SingleUtteranceOmni(),
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        client.put("/api/settings", json={"live_audio": {"mode": "audio_direct"}, "persist": False})
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(f"/api/live-interviews/{session_id}/start", json={"consent_confirmed": True, "expected_revision": 0, "operation_id": str(uuid4())})
        recording_id = _register_live_capture(client, session_id)
        uploaded = client.post(
            f"/api/live-interviews/{session_id}/audio",
            data={
                "mode": "dialogue",
                "speaker": "unknown",
                "language": "zh",
                "recording_id": recording_id,
            },
            files={"file": ("utterance.webm", b"audio-bytes", "audio/webm")},
        )
        assert uploaded.status_code == 200
        live = uploaded.json()["state"]["live_interview"]
        segments = live["segments"]
        assert len(segments) == 1
        assert segments[0]["speaker"] == "candidate"
        assert segments[0]["source"] == "audio_direct"
        assert "用压测验证性能提升" in segments[0]["text"]


def test_mock_interview_unlimited_flow(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'unlimited.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/workflows/candidate-prep",
            json={
                "session_id": session_id,
                "resume_text": "Python",
                "job_description": "Platform",
                "company_name": "Example",
            },
        )
        started = client.post(f"/api/mock-interviews/{session_id}/start").json()
        qid = started["current_question"]["id"]
        # Answer and advance repeatedly; the interview never auto-completes.
        for i in range(4):
            resp = client.post(
                f"/api/mock-interviews/{session_id}/answers",
                json={"question_id": qid, "answer": f"answer {i}"},
            ).json()
            assert resp["mock_session"]["status"] == "active"
            resp = client.post(f"/api/mock-interviews/{session_id}/next").json()
            assert resp["mock_session"]["status"] == "active"
            qid = resp["current_question"]["id"]
        finished = client.post(f"/api/mock-interviews/{session_id}/finish").json()
        assert finished["mock_session"]["status"] == "completed"


def test_mock_interview_retry_endpoint(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'retry-api.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        client.post(
            "/api/workflows/candidate-prep",
            json={
                "session_id": session_id,
                "resume_text": "Python",
                "job_description": "Platform",
                "company_name": "Example",
            },
        )
        started = client.post(f"/api/mock-interviews/{session_id}/start").json()
        qid = started["current_question"]["id"]
        client.post(
            f"/api/mock-interviews/{session_id}/answers",
            json={"question_id": qid, "answer": "first"},
        )
        retried = client.post(
            f"/api/mock-interviews/{session_id}/answers",
            json={"question_id": qid, "answer": "better", "retry": True},
        ).json()
        responses = retried["mock_session"]["responses"]
        assert len(responses) == 1
        assert responses[0]["answer"] == "better"


def test_custom_mock_question_api_adds_and_activates_question(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'custom-question-api.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        response = client.post(
            f"/api/mock-interviews/{session_id}/questions",
            json={
                "question": "你如何处理一次关键项目失败？",
                "competency": "复盘能力",
                "practice_now": True,
            },
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["mock_session"]["status"] == "active"
    assert payload["current_question"]["question"] == "你如何处理一次关键项目失败？"
    assert payload["current_question"]["competency"] == "复盘能力"
    assert payload["current_question"]["source"] == "custom"
    assert payload["current_question"]["answer_framework"]
    assert payload["current_question"]["question_requirements"]
    understanding = payload["current_question"]["understanding"]
    assert understanding["answer_type"] == "behavioral_example"
    assert len(understanding["related_questions"]) == 3
    assert len(understanding["answer_levels"]) == 3
    assert len(understanding["probe_tree"]) == 4
    assert understanding["role_relevance_source"] == "generic"


def test_custom_mock_question_api_validates_length(tmp_path):
    storage = Storage(f"sqlite+aiosqlite:///{tmp_path / 'custom-question-validation.db'}")
    app = create_app(storage=storage, llm_client=WorkflowLLM(), configure_llm=False)
    with TestClient(app) as client:
        session_id = client.post("/api/interviews/sessions", json={}).json()["id"]
        response = client.post(
            f"/api/mock-interviews/{session_id}/questions",
            json={"question": "?"},
        )

    assert response.status_code == 422


def test_app_shutdown_stops_background_before_each_distinct_model_client(tmp_path):
    order = []

    class ClosingClient:
        def __init__(self, name):
            self.name = name

        async def close(self):
            order.append(self.name)

    class ClosingBackground:
        async def close(self):
            order.append("background")

    primary = ClosingClient("primary")
    resume = ClosingClient("resume")
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'shutdown-order.db'}"),
        llm_client=primary,
        resume_llm_client=resume,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        client.app.state.interview_service._background = ClosingBackground()

    assert order == ["background", "primary", "resume"]


def test_app_shutdown_closes_aliased_model_once_and_continues_after_background_error(tmp_path):
    order = []

    class ClosingClient:
        async def close(self):
            order.append("shared-model")

    class FailingBackground:
        async def close(self):
            order.append("background")
            raise RuntimeError("private shutdown detail")

    shared = ClosingClient()
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'shutdown-dedupe.db'}"),
        llm_client=shared,
        resume_llm_client=shared,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    with TestClient(app) as client:
        client.app.state.interview_service._background = FailingBackground()

    assert order == ["background", "shared-model"]


@pytest.mark.asyncio
async def test_app_lifespan_exception_still_shuts_down_tasks_and_clients(tmp_path):
    order = []

    class ClosingClient:
        async def close(self):
            order.append("model")

    class ClosingBackground:
        async def close(self):
            order.append("background")

    model = ClosingClient()
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'shutdown-exception.db'}"),
        llm_client=model,
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )

    with pytest.raises(RuntimeError, match="lifespan consumer failed"):
        async with app.router.lifespan_context(app):
            app.state.interview_service._background = ClosingBackground()
            raise RuntimeError("lifespan consumer failed")

    assert order == ["background", "model"]


@pytest.mark.asyncio
async def test_app_factory_does_not_persist_audio_metadata_before_lifespan(tmp_path):
    settings_path = tmp_path / "settings.json"
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'factory-purity.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=LocalSettingsStore(settings_path),
    )

    assert not settings_path.exists()
    async with app.router.lifespan_context(app):
        assert settings_path.exists()

    persisted = json.loads(settings_path.read_text(encoding="utf-8"))
    assert persisted["_meta"]["audio_settings_etag"]
    assert persisted["_meta"]["audio_settings_revision"] >= 1


@pytest.mark.asyncio
async def test_app_shutdown_finishes_resource_cleanup_before_propagating_cancellation(tmp_path):
    import asyncio

    order = []
    shutdown_started = asyncio.Event()
    allow_shutdown = asyncio.Event()

    class ClosingClient:
        async def close(self):
            order.append("model")

    class BlockingBackground:
        async def close(self):
            order.append("background-start")
            shutdown_started.set()
            await allow_shutdown.wait()
            order.append("background-end")

    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'shutdown-cancel.db'}"),
        llm_client=ClosingClient(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )
    lifespan = app.router.lifespan_context(app)
    await lifespan.__aenter__()
    app.state.interview_service._background = BlockingBackground()

    closing = asyncio.create_task(lifespan.__aexit__(None, None, None))
    await shutdown_started.wait()
    closing.cancel()
    await asyncio.sleep(0)
    allow_shutdown.set()

    with pytest.raises(asyncio.CancelledError):
        await closing
    assert order == ["background-start", "background-end", "model"]


@pytest.mark.asyncio
async def test_reentering_same_app_lifespan_does_not_overwrite_newer_settings(tmp_path):
    settings_path = tmp_path / "settings.json"
    store = LocalSettingsStore(settings_path)
    app = create_app(
        storage=Storage(f"sqlite+aiosqlite:///{tmp_path / 'lifespan-reentry.db'}"),
        llm_client=WorkflowLLM(),
        configure_llm=False,
        settings_store=store,
    )

    async with app.router.lifespan_context(app):
        pass
    persisted = store.load()
    persisted["search"] = {"provider": "none", "search_request_cost_usd": 0.25}
    store.save(persisted)

    async with app.router.lifespan_context(app):
        pass

    assert store.load()["search"]["search_request_cost_usd"] == 0.25


@pytest.mark.asyncio
async def test_app_startup_failure_still_closes_initialized_resources(tmp_path, monkeypatch):
    order = []

    class TrackingStorage(Storage):
        async def close(self):
            order.append("storage")
            await super().close()

    class ClosingClient:
        async def close(self):
            order.append("model")

    async def fail_reconciliation(self):
        raise OSError("private recordings path")

    monkeypatch.setattr(
        InterviewService,
        "cleanup_orphaned_live_audio",
        fail_reconciliation,
    )
    app = create_app(
        storage=TrackingStorage(
            f"sqlite+aiosqlite:///{tmp_path / 'startup-failure-cleanup.db'}"
        ),
        llm_client=ClosingClient(),
        configure_llm=False,
        settings_store=LocalSettingsStore(tmp_path / "settings.json"),
    )

    with pytest.raises(OSError, match="private recordings path"):
        async with app.router.lifespan_context(app):
            pass

    assert order == ["model", "storage"]
